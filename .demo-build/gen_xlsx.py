from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/workspace/demo-output/data.xlsx"

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=11)
ZEBRA_1 = PatternFill(fill_type="solid", fgColor="FFFFFF")
ZEBRA_2 = PatternFill(fill_type="solid", fgColor="F7F9FC")
KPI_FILL = PatternFill(fill_type="solid", fgColor="EAF2FF")
KPI_FONT = Font(name="Arial", bold=True, color="1F2937", size=11)
THIN = Side(style="thin", color="D9DEE7")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_BORDER = Border(left=THIN, right=THIN, top=Side(style="medium", color="AAB4C5"), bottom=THIN)

wb = Workbook()

# ---------- Sheet 1: 格式清单 ----------
ws = wb.active
ws.title = "格式清单"
ws.append(["格式", "类别", "生成工具", "演示文件", "状态"])
rows = [
    ["txt", "文本", "Write", "hello.txt", "已生成"],
    ["md", "文本", "Write", "notes.md", "已生成"],
    ["csv", "文本", "Write", "data.csv", "已生成"],
    ["json", "数据", "Write", "config.json", "已生成"],
    ["yaml", "数据", "Write", "config.yaml", "已生成"],
    ["xml", "数据", "Write", "config.xml", "已生成"],
    ["svg", "矢量图", "Write", "icon.svg", "已生成"],
    ["html", "网页", "Write", "demo.html", "已生成"],
    ["py", "代码", "Write", "demo_code.py", "已生成"],
    ["jpg", "图片", "GenerateImage", "ai_image.jpg", "已生成"],
    ["docx", "Office", "Skill:office-plugin:docx", "report.docx", "已生成"],
    ["xlsx", "Office", "Skill:office-plugin:xlsx", "data.xlsx", "已生成"],
    ["pptx", "Office", "Skill:office-plugin:pptx", "presentation.pptx", "已生成"],
    ["pdf", "Office", "Skill:office-plugin:pdf", "summary.pdf", "已生成"],
    ["zip", "归档", "RunCommand", "demo-output.zip", "已生成"],
]
for r in rows:
    ws.append(r)
last = len(rows) + 1
ws.cell(last + 1, 1, "合计")
ws.cell(last + 1, 2, "=COUNTA(A2:A%d)" % last)
ws.cell(last + 1, 3, "种文件格式")
for c in range(1, 6):
    cell = ws.cell(1, c)
    cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, THIN_BORDER
for r in range(2, last + 1):
    fill = ZEBRA_1 if (r - 2) % 2 == 0 else ZEBRA_2
    for c in range(1, 6):
        cell = ws.cell(r, c)
        cell.fill, cell.font, cell.border = fill, BODY_FONT, THIN_BORDER
for c in range(1, 6):
    cell = ws.cell(last + 1, c)
    cell.fill, cell.font, cell.border = KPI_FILL, KPI_FONT, TOTAL_BORDER
for col, w in zip("ABCDE", [10, 10, 28, 22, 10]):
    ws.column_dimensions[col].width = w

# ---------- Sheet 2: 工具统计 ----------
ws2 = wb.create_sheet("工具统计")
ws2.append(["类别", "工具列表", "数量"])
groups = [
    ["文件与搜索", "Read、Write、Edit、DeleteFile、Glob、Grep、LS、SearchCodebase"],
    ["终端", "RunCommand、CheckCommandStatus、StopCommand"],
    ["网络", "WebSearch、WebFetch"],
    ["图像", "GenerateImage"],
    ["子智能体", "Task"],
    ["规划交互", "TodoWrite、AskUserQuestion、NotifyUser、OpenPreview、browser_waiting_for_user_interaction"],
    ["目标管理", "get_goal、create_goal、update_goal"],
    ["定时任务", "Schedule"],
    ["技能与扩展", "Skill、run_mcp"],
]
for i, (cat, tools) in enumerate(groups, start=2):
    ws2.cell(i, 1, cat)
    ws2.cell(i, 2, tools)
    sep = "、"
    ws2.cell(i, 3, '=LEN(B{0})-LEN(SUBSTITUTE(B{0},"{1}",""))+1'.format(i, sep))
n = len(groups) + 1
ws2.cell(n + 1, 1, "总计")
ws2.cell(n + 1, 3, "=SUM(C2:C%d)" % n)
for c in range(1, 4):
    cell = ws2.cell(1, c)
    cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, THIN_BORDER
for r in range(2, n + 1):
    fill = ZEBRA_1 if (r - 2) % 2 == 0 else ZEBRA_2
    for c in range(1, 4):
        cell = ws2.cell(r, c)
        cell.fill, cell.font, cell.border = fill, BODY_FONT, THIN_BORDER
for c in range(1, 4):
    cell = ws2.cell(n + 1, c)
    cell.fill, cell.font, cell.border = KPI_FILL, KPI_FONT, TOTAL_BORDER
ws2.column_dimensions["A"].width = 14
ws2.column_dimensions["B"].width = 72
ws2.column_dimensions["C"].width = 10

wb.save(OUT)
print("saved:", OUT)
